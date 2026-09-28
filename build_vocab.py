#!/usr/bin/env python3
"""
build_vocab.py — Lớp 1 của Flashcard TOEIC Vocab (ETS).

Đọc 3 nguồn trong folder cha (02_Tu_Vung_ETS):
  1. Boost_Vocab_TOEIC_SUMMARY_ALL.docx  → danh sách từ A–Z (word, def, ipa, tests) + chủ đề
  2. High_Freq_316.xlsx                   → cờ high-frequency + số đề
  3. Boost_Vocab_TOEIC_Reading_TEST{1..10}.docx → câu ví dụ trích từ passage Part 6/7

Ghi:  flashcard/vocab.json
      flashcard/index.html (nếu có) — thay khối <script id="vocab-data" type="application/json">…</script>
      Ghi đè thẳng, không tạo .bak.

Chạy:  python build_vocab.py            (báo cáo ra stdout)
       python build_vocab.py --no-embed (chỉ tạo vocab.json)
"""
from __future__ import annotations

import json
import re
import sys
from collections import Counter, OrderedDict
from pathlib import Path

import docx  # python-docx
import openpyxl
from docx.table import Table
from docx.text.paragraph import Paragraph

HERE = Path(__file__).resolve().parent          # .../02_Tu_Vung_ETS/flashcard
SRC = HERE.parent                                # .../02_Tu_Vung_ETS
SUMMARY = SRC / "Boost_Vocab_TOEIC_SUMMARY_ALL.docx"
HF_XLSX = SRC / "High_Freq_316.xlsx"
TESTS = [SRC / f"Boost_Vocab_TOEIC_Reading_TEST{i}.docx" for i in range(1, 11)]
OUT_JSON = HERE / "vocab.json"
OUT_HTML = HERE / "index.html"

THEMES = [
    "Business / Office", "Travel / Hospitality", "Health / Medical",
    "Education / HR", "Technology / IT", "Retail / Customer Service",
    "Finance / Banking", "Real Estate / Construction", "General",
]
EXAMPLE_MAX = 220


# ---------------------------------------------------------------- helpers
def norm_key(s: str) -> str:
    """Key so khớp: lower, gọn khoảng trắng, giữ nguyên ngoặc (v.)/(n.)."""
    return re.sub(r"\s+", " ", s.strip().lower())


def strip_paren(s: str) -> str:
    return re.sub(r"\s*\([^)]*\)", "", s).strip()


def iter_body(doc):
    """Duyệt body theo thứ tự: ('p', text) hoặc ('t', Table)."""
    for el in doc.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield "p", Paragraph(el, doc).text.strip()
        elif tag == "tbl":
            yield "t", Table(el, doc)


def table_rows(t: Table):
    """Dòng bảng dạng list[str]; python-docx lặp lại dòng merge → caller tự dedup."""
    for r in t.rows:
        yield [c.text.strip() for c in r.cells]


def parse_tests(s: str) -> list[int]:
    return sorted({int(x) for x in re.findall(r"\d+", s or "")})


def word_regex(word: str) -> re.Pattern:
    """Regex tìm từ trong câu: cho phép chia -s/-es/-ed/-d/-ing ở từ cuối; không phân biệt hoa thường."""
    base = strip_paren(word).lower()
    parts = [re.escape(p) for p in base.split()]
    if not parts:
        return re.compile(r"(?!x)x")
    parts[-1] = parts[-1] + r"(?:s|es|ed|d|ing)?"
    return re.compile(r"\b" + r"\s+".join(parts) + r"\b", re.IGNORECASE)


def split_sentences(text: str) -> list[str]:
    text = re.sub(r"\s+", " ", text.replace("\n", " ")).strip()
    # bỏ dòng dẫn "Questions 131–134 refer to the following flyer."
    text = re.sub(r"^Questions\s+\d+[–-]\d+\s+refer to the following[^.]*\.\s*", "", text, flags=re.I)
    text = re.sub(r"\[[^\]]*\]\s*", "", text)          # nhãn [WEB PAGE], [E-MAIL]
    parts = re.split(r"(?<=[.!?])\s+(?=[A-Z\"“(])", text)
    out = []
    for p in parts:
        p = re.sub(r"^(?:[A-Z][A-Z&'-]{2,}\s+)+(?=[A-Z][a-z])", "", p.strip())  # tiêu đề HOA đầu câu: "COUPON If you…"
        if len(p) >= 25:
            out.append(p)
    return out


def pick_example(cands: list[tuple[int, str]]) -> str:
    """cands: (test_no, sentence). Ưu tiên câu ≤ EXAMPLE_MAX, độ dài gần 110; hết thì cắt câu ngắn nhất."""
    if not cands:
        return ""
    ok = [s for _, s in cands if len(s) <= EXAMPLE_MAX]
    if ok:
        return min(ok, key=lambda s: abs(len(s) - 110))
    s = min((s for _, s in cands), key=len)
    return s[: EXAMPLE_MAX - 1].rsplit(" ", 1)[0] + "…"


# ---------------------------------------------------------------- 1. SUMMARY A–Z + theme
def read_summary():
    doc = docx.Document(SUMMARY)
    entries: "OrderedDict[str, dict]" = OrderedDict()
    theme_of: dict[str, str] = {}
    section = None           # "az" | "theme"
    cur_theme = None
    dup_rows = 0

    for kind, val in iter_body(doc):
        if kind == "p":
            if val.startswith("PHẦN 1"):
                section = "az"
            elif val.startswith("PHẦN 2"):
                section = "theme"
            elif section == "theme":
                for th in THEMES:
                    if val.startswith(th):
                        cur_theme = th
            continue

        rows = list(table_rows(val))
        if not rows or rows[0][0].lower() != "word":
            continue
        if section == "az":
            for r in rows[1:]:
                if len(r) < 4 or not r[0]:
                    continue
                k = norm_key(r[0])
                tests = parse_tests(r[3])
                if k in entries:
                    dup_rows += 1
                    e = entries[k]
                    e["tests"] = sorted(set(e["tests"]) | set(tests))
                    if not e["def"] and r[1]:
                        e["def"] = r[1]
                    if not e["ipa"] and r[2]:
                        e["ipa"] = r[2]
                    continue
                entries[k] = {
                    "id": k, "word": r[0], "ipa": r[2],
                    "def": r[1], "tests": tests,
                }
        elif section == "theme" and cur_theme:
            for r in rows[1:]:
                if r and r[0]:
                    theme_of.setdefault(norm_key(r[0]), cur_theme)

    for k, e in entries.items():
        e["theme"] = theme_of.get(k, "General")
    return entries, dup_rows, theme_of


# ---------------------------------------------------------------- 2. High-frequency xlsx
def read_hf():
    wb = openpyxl.load_workbook(HF_XLSX, read_only=True)
    ws = wb.active
    hf: dict[str, int] = {}
    for row in ws.iter_rows(min_row=1, values_only=True):
        if not row or not isinstance(row[1], str):
            continue
        w = row[1].strip()
        if not w or w.lower() == "word" or w.startswith("="):
            continue
        try:
            n = int(row[4])
        except (TypeError, ValueError):
            n = 0
        hf[norm_key(w)] = n
    return hf


# ---------------------------------------------------------------- 3. Examples from TEST docx
KW_LINE = re.compile(r"^(?P<w>[^=\n]+?)\s*=\s*(?P<rest>.+)$")


def read_examples(entries):
    cands: dict[str, list[tuple[int, str]]] = {k: [] for k in entries}
    kw_total = kw_unmatched = 0
    unmatched_samples: Counter = Counter()

    # cache regex theo key
    rx_cache: dict[str, re.Pattern] = {}

    for test_no, path in enumerate(TESTS, start=1):
        if not path.exists():
            print(f"  ! thiếu {path.name}", file=sys.stderr)
            continue
        doc = docx.Document(path)
        for kind, val in iter_body(doc):
            if kind != "t" or len(val.columns) != 2:
                continue
            rows = list(table_rows(val))
            if not rows:
                continue
            passage, kwcol = rows[0][0], rows[0][1]
            sentences = split_sentences(passage)
            if not sentences:
                continue
            for line in kwcol.splitlines():
                line = line.strip()
                m = KW_LINE.match(line)
                if not m:
                    continue
                kw_total += 1
                raw = m.group("w").strip()
                k = norm_key(raw)
                if k not in entries:
                    k2 = norm_key(strip_paren(raw))
                    # thử khớp id có ngoặc trong summary với từ trần trong test
                    if k2 in entries:
                        k = k2
                    else:
                        alt = [e for e in entries if strip_paren(e) == k2]
                        if len(alt) == 1:
                            k = alt[0]
                        else:
                            kw_unmatched += 1
                            unmatched_samples[raw] += 1
                            continue
                rx = rx_cache.get(k)
                if rx is None:
                    rx = rx_cache[k] = word_regex(entries[k]["word"])
                for s in sentences:
                    if rx.search(s):
                        cands[k].append((test_no, s))
                        break
    return cands, kw_total, kw_unmatched, unmatched_samples


# ---------------------------------------------------------------- embed
def embed(json_text: str) -> bool:
    if not OUT_HTML.exists():
        return False
    html = OUT_HTML.read_text(encoding="utf-8")
    pat = re.compile(r'(<script id="vocab-data" type="application/json">)(.*?)(</script>)', re.S)
    if not pat.search(html):
        print(f"  ! {OUT_HTML.name} không có khối vocab-data — bỏ qua embed")
        return False
    safe = json_text.replace("</", "<\\/")
    html = pat.sub(lambda m: m.group(1) + safe + m.group(3), html, count=1)
    OUT_HTML.write_text(html, encoding="utf-8")
    return True


def sync_version() -> str:
    """Đọc APP_VER trong index.html, ghi version.json cùng folder (app so sánh để tự tải lại bản mới)."""
    if not OUT_HTML.exists():
        return "bỏ qua"
    m = re.search(r'const APP_VER = "([^"]+)"', OUT_HTML.read_text(encoding="utf-8"))
    if not m:
        return "không thấy APP_VER"
    (HERE / "version.json").write_text(json.dumps({"ver": m.group(1)}), encoding="utf-8")
    return m.group(1)


# ---------------------------------------------------------------- main
def main(argv):
    no_embed = "--no-embed" in argv
    entries, dup_rows, theme_of = read_summary()
    hf = read_hf()
    cands, kw_total, kw_unmatched, unmatched = read_examples(entries)

    hf_hit = 0
    for k, e in entries.items():
        if k in hf:
            e["hf"] = True
            e["count"] = hf[k] or len(e["tests"])
            hf_hit += 1
        else:
            e["hf"] = False
            e["count"] = len(e["tests"])
        e["example"] = pick_example(cands.get(k, []))

    hf_miss = [w for w in hf if w not in entries]
    out = list(entries.values())
    OUT_JSON.write_text(json.dumps(out, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    compact = json.dumps(out, ensure_ascii=False, separators=(",", ":"))
    embedded = False if no_embed else embed(compact)

    n = len(out)
    n_ex = sum(1 for e in out if e["example"])
    n_ex_hf = sum(1 for e in out if e["example"] and e["hf"])
    themes = Counter(e["theme"] for e in out)
    print("=== build_vocab report ===")
    print(f"Entries            : {n}  (dòng trùng đã gộp: {dup_rows})")
    print(f"HF-316 khớp        : {hf_hit}/{len(hf)}  | không khớp: {len(hf_miss)} {hf_miss[:10]}")
    print(f"Có câu ví dụ       : {n_ex}/{n} ({n_ex*100//n}%)  | trong HF: {n_ex_hf}/{hf_hit}")
    print(f"Keyword trong TEST : {kw_total}  | không khớp summary: {kw_unmatched}")
    if unmatched:
        print("   mẫu không khớp  :", ", ".join(f"{w}" for w, _ in unmatched.most_common(12)))
    print("Theme              :", ", ".join(f"{t}={c}" for t, c in themes.most_common()))
    print(f"vocab.json         : {OUT_JSON.stat().st_size/1024:.0f} KB → {OUT_JSON}")
    print(f"embed index.html   : {'OK' if embedded else 'bỏ qua'}")
    print(f"version.json       : {sync_version()}")


if __name__ == "__main__":
    main(sys.argv[1:])
